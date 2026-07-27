#!/usr/bin/env python3
"""Extract text and structural warnings from resumes and career documents.

Two jobs in one module:

1. Text extraction (PDF, DOCX, TXT, MD) so the evidence connectors have
   something to read.
2. Structural inspection, because how a resume is BUILT determines whether an
   ATS can read it. A two-column layout renders beautifully for a human and
   parses into interleaved garbage for a machine. That failure is invisible
   unless you look at the structure, not the text.

Usage:
    python scripts/ingest.py <path> [--json] [--structure]
    python scripts/ingest.py <dir> --walk [--json]
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import unicodedata
from dataclasses import dataclass, field, asdict
from pathlib import Path

TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".rst"}
SUPPORTED = TEXT_SUFFIXES | {".pdf", ".docx"}


@dataclass
class Doc:
    path: str
    kind: str                     # pdf | docx | text
    text: str
    pages: int = 0
    words: int = 0
    warnings: list[str] = field(default_factory=list)
    structure: dict = field(default_factory=dict)
    error: str | None = None


def _missing(pkg: str, why: str) -> str:
    return f"{pkg} not installed, so {why}. Run: pip install -r requirements.txt"


# Control characters that are never legitimate in extracted document text.
# Tab, newline, and carriage return are kept; everything else in the C0 and C1
# ranges is stripped.
_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]")


def sanitize(text: str) -> str:
    """Strip control characters from extracted text.

    PDF text layers regularly contain stray null bytes and other control
    characters, especially from LaTeX-generated files. A single null byte makes
    grep treat the file as binary, can truncate a C-backed string reader, and
    is meaningless to every consumer downstream. Strip once, here, so nothing
    below this layer has to think about it.
    """
    return _CONTROL.sub("", text)


# ------------------------------------------------------------------ PDF

def _has_gutter(page, min_width: float = 18.0, min_run: int = 3) -> bool:
    """True when a vertical whitespace band splits text into parallel columns.

    Two things this must get right, because the naive version of this check
    gets both wrong.

    It must NOT fire on right-aligned dates. Every resume right-aligns dates,
    which puts 30 to 40 percent of words in the right half of a perfectly
    ordinary single-column page. The tell is that right-aligned content all
    ENDS at the right margin and STARTS wherever the preceding text happened
    to stop, so its start positions are scattered.

    It must still fire on a localized column block, such as a three-column
    skills list occupying four rows of an otherwise single-column resume.
    That block genuinely interleaves when parsed, so requiring a gutter down
    the whole page misses a real defect.

    So: find runs of consecutive rows sharing a gutter, then require the right
    side of those rows to start at a consistent x, which is what a column does
    and right-alignment does not.
    """
    words = page.extract_words() or []
    if len(words) < 30:
        return False

    buckets: dict[int, list] = {}
    for w in words:
        buckets.setdefault(int(w["top"] // 6), []).append(w)
    keys = sorted(k for k, v in buckets.items() if len(v) >= 2)
    if len(keys) < min_run:
        return False

    right_margin = max(w["x1"] for w in words)

    x = page.width * 0.15
    while x < page.width * 0.85:
        run: list[int] = []
        for k in keys:
            ws = buckets[k]
            before = any(w["x1"] <= x for w in ws)
            after = any(w["x0"] >= x + min_width for w in ws)
            inside = any(w["x0"] < x + min_width and w["x1"] > x for w in ws)
            if before and after and not inside:
                run.append(k)
                continue
            if len(run) >= min_run and _run_is_columnar(buckets, run, x, min_width, right_margin):
                return True
            run = []
        if len(run) >= min_run and _run_is_columnar(buckets, run, x, min_width, right_margin):
            return True
        x += 4.0
    return False


def _run_is_columnar(buckets, run, x, min_width, right_margin) -> bool:
    """Distinguish a real column from a run of right-aligned line endings."""
    starts, ends, cell_words = [], [], []
    for k in run:
        right = [w for w in buckets[k] if w["x0"] >= x + min_width]
        if not right:
            return False
        starts.append(min(w["x0"] for w in right))
        ends.append(max(w["x1"] for w in right))
        cell_words.append(len(right))

    spread = max(starts) - min(starts)

    # Right-aligned content ends flush against the right margin on every row.
    if sum(1 for e in ends if abs(e - right_margin) < 3.0) >= len(run) - 1 and spread > 20.0:
        return False

    # A column carries real content. A right-aligned date or a repeated
    # "Certification" label is one to three short tokens, and three such rows
    # stacked up are geometrically identical to a column while parsing
    # perfectly well, because each row is a self-contained unit.
    if sum(cell_words) / len(cell_words) < 2.0:
        return False

    # A real column starts at the same x on every row.
    return spread <= 20.0


def read_pdf(path: Path) -> Doc:
    d = Doc(path=str(path), kind="pdf", text="")
    try:
        import pdfplumber
    except ImportError:
        d.error = _missing("pdfplumber", "PDF text could not be extracted")
        return d

    chunks: list[str] = []
    col_suspect = 0
    table_pages = 0

    no_space_glyphs = False

    try:
        with pdfplumber.open(path) as pdf:
            d.pages = len(pdf.pages)

            # Does the file contain real space characters at all? Many LaTeX
            # resume templates emit none: every gap between words is glyph
            # positioning, not a space. The PDF looks perfect and the text
            # layer is one continuous string.
            total_chars = sum(len(p.chars) for p in pdf.pages)
            space_chars = sum(1 for p in pdf.pages for c in p.chars if c["text"] == " ")
            if total_chars > 200 and space_chars / total_chars < 0.02:
                no_space_glyphs = True

            for page in pdf.pages:
                # pdfplumber defaults to x_tolerance=3.0, which glues words
                # together when the font's inter-word gap is narrow. Tighten it
                # for files with no space glyphs so we recover real words to
                # analyze, rather than analyzing mush of our own making.
                chunks.append(page.extract_text(x_tolerance=1.5 if no_space_glyphs else 3.0) or "")

                # Two-column detection by finding a vertical gutter.
                #
                # Counting words past the page midpoint does NOT work. Every
                # resume right-aligns its dates, which puts 30% of words in the
                # right half of a perfectly ordinary single-column document.
                # What actually distinguishes a two-column layout is a vertical
                # band of whitespace that persists down the page while text
                # exists on both sides of it.
                if _has_gutter(page):
                    col_suspect += 1

                if page.find_tables():
                    table_pages += 1
    except Exception as e:  # pdfplumber raises a wide variety on damaged files
        d.error = f"could not read PDF: {e}"
        return d

    d.text = "\n".join(chunks).strip()
    d.structure = {"columns_suspected": col_suspect, "table_pages": table_pages,
                   "no_space_glyphs": no_space_glyphs}

    if no_space_glyphs:
        d.warnings.append(
            "the PDF contains no space characters. Every gap between words is "
            "glyph positioning, not a space. Sophisticated parsers reconstruct "
            "the words from coordinates; simpler ones return one continuous "
            "string, so 'in PyTorch' becomes 'inPyTorch' and a keyword search "
            "for PyTorch finds nothing. Common with LaTeX resume templates. "
            "Fix by exporting from a different engine, or ship DOCX instead."
        )

    if col_suspect:
        d.warnings.append(
            f"multi-column layout detected on {col_suspect} page(s). "
            "ATS parsers read left to right across the whole line and will "
            "interleave your columns into nonsense. Use a single column."
        )
    if table_pages:
        d.warnings.append(
            f"table structures on {table_pages} page(s). Many ATS parsers drop "
            "table contents entirely. Move anything important out of tables."
        )
    if d.pages > 2:
        d.warnings.append(f"{d.pages} pages. Cut to 1 page unless you have 10+ years of experience.")
    if not d.text.strip():
        d.warnings.append(
            "no extractable text. This PDF is an image or has broken font "
            "encoding. An ATS sees a blank document. Re-export from the source."
        )

    # Unmapped glyphs. LaTeX resume templates that use icon fonts (FontAwesome
    # for the phone, email, GitHub icons) very often ship without a Unicode
    # mapping. The PDF renders a pretty icon for a human and emits literal
    # "(cid:211)" to any parser. This sits right next to the contact details,
    # which is the worst possible place for garbage.
    cids = re.findall(r"\(cid:\d+\)", d.text)
    if cids:
        d.structure["unmapped_glyphs"] = len(cids)
        d.warnings.append(
            f"{len(cids)} unmapped font glyph(s) such as {cids[0]} in the extracted "
            "text. These are icon-font characters with no Unicode mapping. An ATS "
            "reads them as literal garbage. Replace decorative icons with plain "
            "text labels."
        )

    # Stray modifier letters and private-use characters, same root cause.
    strays = [c for c in set(d.text) if unicodedata.category(c) in ("Lm", "Sk", "Co")]
    if strays:
        d.warnings.append(
            "stray decorative characters in the text layer: "
            + " ".join(repr(c) for c in sorted(strays)[:6])
            + ". Same icon-font problem. Remove them."
        )
    return d


# ------------------------------------------------------------------ DOCX

def read_docx(path: Path) -> Doc:
    d = Doc(path=str(path), kind="docx", text="")
    try:
        import docx  # python-docx
    except ImportError:
        d.error = _missing("python-docx", "DOCX text could not be extracted")
        return d

    try:
        doc = docx.Document(str(path))
    except Exception as e:
        d.error = f"could not read DOCX: {e}"
        return d

    parts = [p.text for p in doc.paragraphs]

    n_tables = len(doc.tables)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                parts.append(cell.text)

    n_header = sum(len(s.header.paragraphs) for s in doc.sections)
    n_footer = sum(len(s.footer.paragraphs) for s in doc.sections)
    header_text = " ".join(
        p.text for s in doc.sections for p in s.header.paragraphs
    ).strip()

    d.text = "\n".join(x for x in parts if x.strip()).strip()
    d.structure = {"tables": n_tables, "header_paras": n_header, "footer_paras": n_footer}

    if n_tables:
        d.warnings.append(
            f"{n_tables} table(s). Some ATS parsers skip table contents. "
            "If your contact block or skills live in a table, they may vanish."
        )
    if header_text:
        d.warnings.append(
            "content in the page header. Several major ATS platforms ignore "
            "headers and footers completely. If your name and email are up "
            "there, the system has no idea who applied."
        )
    return d


# ------------------------------------------------------------------ plain

def read_text(path: Path) -> Doc:
    d = Doc(path=str(path), kind="text", text="")
    try:
        d.text = path.read_text(encoding="utf-8", errors="replace")
    except Exception as e:
        d.error = f"could not read: {e}"
    return d


# ------------------------------------------------------------------ api

def read(path: Path) -> Doc:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        d = read_pdf(path)
    elif suffix == ".docx":
        d = read_docx(path)
    elif suffix in TEXT_SUFFIXES:
        d = read_text(path)
    elif suffix == ".doc":
        d = Doc(path=str(path), kind="doc", text="",
                error="legacy .doc is not supported. Open it and save as .docx.")
    else:
        d = Doc(path=str(path), kind="unknown", text="",
                error=f"unsupported file type: {suffix or 'no extension'}")

    d.text = sanitize(d.text)
    d.words = len(re.findall(r"\b[\w'-]+\b", d.text))
    return d


def walk(root: Path, skip: set[str] | None = None) -> list[Doc]:
    """Read every supported document under root.

    `skip` holds directory names to refuse outright. Callers pass the
    financial and immigration folders here.
    """
    skip = skip or set()
    docs: list[Doc] = []
    for p in sorted(root.rglob("*")):
        if not p.is_file():
            continue
        if p.suffix.lower() not in SUPPORTED:
            continue
        if any(part in skip for part in p.parts):
            continue
        if p.name.startswith((".", "~$")):
            continue
        docs.append(read(p))
    return docs


def main() -> int:
    ap = argparse.ArgumentParser(description="Extract text and structure from documents.")
    ap.add_argument("path")
    ap.add_argument("--walk", action="store_true", help="recurse a directory")
    ap.add_argument("--json", action="store_true")
    ap.add_argument("--structure", action="store_true", help="warnings only, no body text")
    args = ap.parse_args()

    root = Path(args.path).expanduser()
    if not root.exists():
        print(f"no such path: {root}", file=sys.stderr)
        return 2

    docs = walk(root) if args.walk else [read(root)]

    if args.json:
        payload = [asdict(d) for d in docs]
        if args.structure:
            for p in payload:
                p.pop("text", None)
        print(json.dumps(payload, indent=2))
        return 0

    for d in docs:
        print(f"\n=== {d.path}")
        if d.error:
            print(f"  ERROR  {d.error}")
            continue
        print(f"  {d.kind}, {d.words} words" + (f", {d.pages} pages" if d.pages else ""))
        for w in d.warnings:
            print(f"  WARN   {w}")
        if not args.structure and d.text:
            preview = d.text[:600].replace("\n", "\n  ")
            print(f"\n  {preview}{' ...' if len(d.text) > 600 else ''}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
